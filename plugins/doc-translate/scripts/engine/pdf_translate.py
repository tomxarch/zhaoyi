"""照译 第 11–13 天：翻译 PDF。

三条路线：
- babeldoc：调用开源的 BabelDOC，保留原 PDF 排版，输出 PDF。扫描件可以打开 OCR。
  BabelDOC 是 AGPL-3.0 协议，所以这里不把它打包进照译，而是调用单独安装的 babeldoc 命令。
- vision：把每页渲染成图片交给视觉模型，直接看图翻译，输出 Word。
  适合很乱的扫描件（污渍、印章、手写），排版会简化成标题、段落、列表和表格。
- overlay：同样看图，但把译文盖在原页面的文字位置上，输出 PDF（像 DeepL 那样），见 pdf_overlay.py。
  扫描件、没装 BabelDOC 时默认走这条。
"""
import base64
import io
import re
import shutil
import subprocess
import time
from pathlib import Path

from providers import LANGUAGES, VISION_MODELS, with_retry

# 支持看图的模型；None 表示直接用翻译时选的模型（Gemini 的文字模型本身就能看图）

VISION_PROMPT = ("这是一份文档的第 {page} 页。把页面上的全部文字翻译成{lang}，"
                 "按原来的结构输出 Markdown：标题用 #，列表用 -，表格用 Markdown 表格。"
                 "图片、logo、示意图、照片在它原来的位置单独写一行 [[图片 ymin,xmin,ymax,xmax]]，"
                 "坐标按 0 到 1000 归一化，框住整张图（图里的文字不用翻）。"
                 "印章、手写、看不清的地方用【】注明，【】里的说明也要用{lang}写。"
                 "页码、页眉页脚不要输出；表格单元格里不要用 <br> 等 HTML 标签；"
                 "不要用 LaTeX，符号直接写成 φ、×、±、≤、°、Ω 这样的字符。"
                 "只输出译文，不要解释。")


def translate_pdf(src, translator, mode="auto", ocr=False, glossary_file=None,
                  bilingual=False, on_progress=None, on_log=print):
    """mode：auto（默认）、babeldoc、overlay（盖译文，输出 PDF）、vision（看图翻译，输出 Word）。
    auto：有文字层且装了 BabelDOC 就用 BabelDOC，否则（扫描件或没装）用 overlay。"""
    if mode == "auto":
        mode = ("babeldoc" if translator.kind == "openai" and not getattr(translator, "drafter", None) and find_babeldoc() and not ocr and has_text_layer(src)
                else "overlay")  # BabelDOC 只能接大模型，也用不上机器初译
    if mode == "vision":
        return [translate_pdf_vision(src, translator, on_progress=on_progress, on_log=on_log)]
    if mode == "overlay":
        from pdf_overlay import rapid_ocr, translate_pdf_overlay
        if translator.provider in VISION_MODELS or not rapid_ocr():
            model, client = vision_client(translator)
        else:  # DeepSeek 这类不能看图的模型：认字全交给本机 OCR，模型只管翻译
            model, client = None, None
        on_log("PDF 先认出文字，再把译文盖在原来的位置，输出 PDF。")
        return [translate_pdf_overlay(src, translator, model, client, on_progress=on_progress, on_log=on_log)]
    return translate_pdf_babeldoc(src, translator, ocr=ocr, glossary_file=glossary_file,
                                  bilingual=bilingual, on_log=on_log)


def has_text_layer(src, pages=3):
    """前几页能抽出足够多的文字，就算有文字层（不是扫描件）。"""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(src))
    chars = sum(len(pdf[i].get_textpage().get_text_range().strip()) for i in range(min(pages, len(pdf))))
    return chars > 50


def vision_client(translator):
    if translator.provider not in VISION_MODELS:
        raise ValueError(f"{translator.label} 不能看图。输出 PDF（保留原版面）可以用它，"
                         "要输出可编辑的 Word（看图翻译），请换 Gemini、通义、智谱、Kimi、豆包、硅基流动或 Claude。")
    from openai import OpenAI
    client = OpenAI(api_key=translator.key, base_url=translator.base_url, max_retries=0, timeout=120)
    return VISION_MODELS[translator.provider] or translator.model, client


# ---------- 路线一：BabelDOC ----------

def find_babeldoc():
    exe = shutil.which("babeldoc")
    if exe:
        return exe
    # uv tool 默认装在这里
    for candidate in (Path.home() / ".local" / "bin" / "babeldoc.exe", Path.home() / ".local" / "bin" / "babeldoc"):
        if candidate.exists():
            return str(candidate)
    return None


def translate_pdf_babeldoc(src, translator, ocr=False, glossary_file=None, bilingual=False, on_log=print):
    if translator.kind != "openai":
        raise ValueError("PDF 翻译需要大模型（Gemini、DeepSeek 或通义），DeepL 和 test 用不了。")
    exe = find_babeldoc()
    if not exe:
        raise RuntimeError("还没有安装 BabelDOC。请按 README 的“PDF 翻译”一节安装。")

    src = Path(src).resolve()
    out_dir = src.parent
    lang_out = LANGUAGES[translator.target_lang][1]
    lang_in = "en" if lang_out == "zh" else "zh"
    cmd = [exe, "--files", str(src), "--output", str(out_dir),
           "--lang-in", lang_in, "--lang-out", lang_out,
           "--openai", "--openai-model", translator.model,
           "--openai-base-url", translator.base_url, "--openai-api-key", translator.key,
           "--watermark-output-mode", "no_watermark", "--qps", "4"]
    if not bilingual:
        cmd.append("--no-dual")
    if ocr:
        cmd.append("--ocr-workaround")
    else:
        cmd.append("--auto-enable-ocr-workaround")
    if glossary_file:
        cmd += ["--glossary-files", str(Path(glossary_file).resolve())]
    if translator.scene:
        cmd += ["--custom-system-prompt", translator.scene]

    started = time.time()
    on_log("BabelDOC 开始处理，第一次运行要下载版面识别模型，会慢一些……")
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            text=True, encoding="utf-8", errors="replace",
                            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    tail = []
    for line in proc.stdout:
        line = line.strip()
        if line:
            tail = (tail + [line])[-20:]
            if any(w in line.lower() for w in ("error", "translat", "finish", "ocr")):
                on_log("  " + line.replace(translator.key, "***")[:200])
    if proc.wait() != 0:
        raise RuntimeError("BabelDOC 出错了，最后几行：\n" + "\n".join(tail).replace(translator.key, "***"))

    outputs = [p for p in out_dir.glob(f"{src.stem}*.pdf") if p != src and p.stat().st_mtime >= started]
    if not outputs:
        raise RuntimeError("BabelDOC 跑完了，但没找到输出文件。")
    return sorted(outputs)


# ---------- 路线二：视觉模型看图翻译 ----------

def render_pages(src, dpi=150):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(src))
    for i in range(len(pdf)):
        image = pdf[i].render(scale=dpi / 72).to_pil().convert("RGB")
        buf = io.BytesIO()
        image.save(buf, format="JPEG", quality=85)
        yield i + 1, len(pdf), base64.b64encode(buf.getvalue()).decode(), image, pdf[i].get_size()


def translate_pdf_vision(src, translator, on_progress=None, on_log=print):
    model, client = vision_client(translator)

    pages, shots, size = [], [], None
    for page, total, image, shot, page_size in render_pages(src):
        shots.append(shot)
        size = size or page_size
        prompt = translator.scene + VISION_PROMPT.format(page=page, lang=translator.target_lang)
        if translator.glossary:
            prompt += "术语必须这样翻译：" + "；".join(f"{s} → {t}" for s, t in translator.glossary.items())

        def call():
            response = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{image}"}},
                ]}],
                temperature=0.2,
            )
            return response.choices[0].message.content

        pages.append(with_retry(call))
        if on_progress:
            on_progress(page, total)

    src = Path(src)
    dst = src.with_name(f"{src.stem}_看图译文.docx")
    markdown_to_docx(pages, dst, shots, size)
    on_log("看图翻译会重新排版：图片从原页抠出放回，标题、段落、表格重新排，版式需要你再核对一下。")
    return dst


IMAGE_LINE = re.compile(r"\[\[\s*(?:图片|image|figure)?\s*:?\s*\[?([\d.\s,]+)\]?\s*\]\]", re.I)


def add_picture(doc, shot, box, usable_width):
    """按模型给的位置从原页面抠图，宽度按它在原页上占的比例。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    try:
        ymin, xmin, ymax, xmax = (max(0.0, min(1000.0, float(v))) for v in box.replace(" ", "").split(",")[:4])
    except ValueError:
        return
    if ymax - ymin < 5 or xmax - xmin < 5:
        return
    w, h = shot.size
    pad = 6
    crop = shot.crop((max(int(xmin / 1000 * w) - pad, 0), max(int(ymin / 1000 * h) - pad, 0),
                      min(int(xmax / 1000 * w) + pad, w), min(int(ymax / 1000 * h) + pad, h)))
    buf = io.BytesIO()
    crop.save(buf, format="PNG")
    buf.seek(0)
    paragraph = doc.add_paragraph()
    center = abs((xmin + xmax) / 2 - 500) < 80
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(buf, width=int(usable_width * min((xmax - xmin) / 1000 * 1.05, 1)))


def markdown_to_docx(pages, dst, shots=None, page_size=None):
    """pages：每页的 Markdown；shots：每页原图（用来抠图片）；page_size：原 PDF 页面大小（磅）。"""
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    section = doc.sections[0]
    if page_size:
        section.page_width, section.page_height = Pt(page_size[0]), Pt(page_size[1])
        section.left_margin = section.right_margin = Pt(page_size[0] * 0.08)
    usable = section.page_width - section.left_margin - section.right_margin
    for n, text in enumerate(pages):
        if n:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        text = re.sub(r"^```\w*\n|\n```$", "", text.strip())
        lines = text.splitlines()
        # 模型偶尔会把页码当正文输出，去掉页首页尾单独的数字
        while lines and re.fullmatch(r"\s*-?\s*\d{1,3}\s*-?\s*", lines[-1]):
            lines.pop()
        while lines and re.fullmatch(r"\s*-?\s*\d{1,3}\s*-?\s*", lines[0]):
            lines.pop(0)
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            picture = IMAGE_LINE.search(line)
            if picture:
                if shots and n < len(shots):
                    add_picture(doc, shots[n], picture.group(1), usable)
                i += 1
                continue
            if line.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                if rows:
                    cols = max(len(r) for r in rows)
                    table = doc.add_table(rows=len(rows), cols=cols, style="Table Grid")
                    for r, cells in enumerate(rows):
                        for c, cell in enumerate(cells):
                            add_inline(table.cell(r, c).paragraphs[0], cell)
                continue
            heading = re.match(r"(#{1,4})\s+(.*)", line)
            if heading:
                add_inline(doc.add_heading(level=len(heading.group(1))), heading.group(2))
            elif re.match(r"\s*[-*]\s+", line):
                add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"\s*[-*]\s+", "", line, count=1))
            elif re.match(r"\s*\d+[.)]\s+", line):
                add_inline(doc.add_paragraph(style="List Number"), re.sub(r"\s*\d+[.)]\s+", "", line, count=1))
            elif line.strip():
                add_inline(doc.add_paragraph(), line)
            i += 1
    doc.save(dst)


LATEX_SYMBOLS = {"phi": "φ", "varphi": "φ", "times": "×", "pm": "±", "le": "≤", "leq": "≤", "ge": "≥",
                 "geq": "≥", "circ": "°", "degree": "°", "Omega": "Ω", "omega": "ω", "mu": "μ", "alpha": "α",
                 "beta": "β", "Delta": "Δ", "delta": "δ", "approx": "≈", "sim": "~", "cdot": "·", "%": "%"}


def strip_latex(text):
    """模型有时把符号写成 LaTeX，比如 $\\phi 10$、$\\times 40$，换回普通字符。"""
    def plain(m):
        body = re.sub(r"\\([A-Za-z]+|%)", lambda k: LATEX_SYMBOLS.get(k.group(1), k.group(1)), m.group(1))
        body = re.sub(r"\^\{?\\?circ\}?|\^\{?°\}?", "°", body)
        return re.sub(r"[{}]|\s+(?=[×±·°])|(?<=[×±φ·])\s+", "", body).replace("^", "").replace("_", "")
    return re.sub(r"\$([^$\n]{1,60})\$", plain, text)


def add_inline(paragraph, text):
    """支持 **加粗** 和 <br> 换行，其余原样写入。"""
    text = strip_latex(text)
    text = re.sub(r"(?:<|&lt;)br\s*/?(?:>|&gt;)", "\n", text, flags=re.I)
    for k, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        for n, line in enumerate(piece.split("\n")):
            if n:
                paragraph.add_run().add_break()
            if line:
                paragraph.add_run(line).bold = bool(k % 2)
